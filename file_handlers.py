import base64# file_handlers.py
import base64
import PyPDF2
import docx

def encode_image(image_file):
    return base64.b64encode(image_file.read()).decode('utf-8')

def extract_text_from_pdf(file):
    try:
        pdf_reader = PyPDF2.PdfReader(file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
        return text[:3000]
    except Exception:
        return "无法读取PDF内容"

def extract_text_from_docx(file):
    try:
        doc = docx.Document(file)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text[:3000]
    except Exception:
        return "无法读取Word文档内容"

def extract_text_from_txt(file):
    try:
        return file.read().decode('utf-8')[:3000]
    except Exception:
        return file.read().decode('gbk')[:3000]
import PyPDF2
import docx

def encode_image(image_file):
    return base64.b64encode(image_file.read()).decode('utf-8')

def extract_text_from_pdf(file):
    try:
        pdf_reader = PyPDF2.PdfReader(file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
        return text[:3000]
    except:
        return "无法读取PDF内容"

def extract_text_from_docx(file):
    try:
        doc = docx.Document(file)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text[:3000]
    except:
        return "无法读取Word文档内容"

def extract_text_from_txt(file):
    try:
        return file.read().decode('utf-8')[:3000]
    except:
        return file.read().decode('gbk')[:3000]
